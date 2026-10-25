#!/usr/bin/env python3
"""Build news.json for the 「ニュースまとめ」 view.

The Excel workbook is the master for which items exist and how they are worded.
Optional Google Sheet CSV exports supply the full source-URL lists (the Excel
cell keeps only one hyperlink). Daily updates should regenerate news.json only;
index.html does not need to change.

Requires: Python 3, openpyxl (`pip install openpyxl`).
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import OrderedDict
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

INDUSTRY_SHEET = "旅行業界ニュース"
OTA_SHEET = "海外OTAニュース"

SHEETS = OrderedDict(
    [
        (
            INDUSTRY_SHEET,
            {
                "feed": "industry",
                "feedLabel": "旅行業界",
                "body": "内容",
                "note": "話題の切り口",
                "noteLabel": "話題の切り口",
                "sources": "出典",
            },
        ),
        (
            OTA_SHEET,
            {
                "feed": "ota",
                "feedLabel": "海外OTA",
                "body": "要点",
                "note": "実務メモ",
                "noteLabel": "実務メモ",
                "sources": "ソース",
            },
        ),
    ]
)

# フィルタの並び。要点はチップにせず、日付ごとの要約としてカード先頭に出す。
# B2B・卸（Expedia、Hotelbeds、サイトコントローラー、配信規格など）は海外OTAに含める。
# チップは「すべて」の右に 新規開業ホテル、その次に 新規フライト情報。
# 海外OTAと国内OTA・旅行会社は項目のテーマ名として残し、「すべて」に出す。チップには出さない。
# 行政・インバウンド施策とその他は分類はするが、チップにも「すべて」にも出さない。
THEMES = [
    ("overseas", "海外OTA"),
    ("domestic", "国内OTA・旅行会社"),
    ("openings", "新規開業ホテル"),
    ("flights", "新規フライト情報"),
]
CHIP_THEMES = [
    ("openings", "新規開業ホテル"),
    ("flights", "新規フライト情報"),
]
OMITTED_THEMES = (
    ("policy", "行政・インバウンド施策"),
    ("other", "その他"),
)
THEME_LABEL = dict(THEMES + list(OMITTED_THEMES))
OMITTED_IDS = {theme_id for theme_id, _label in OMITTED_THEMES}

WEEKDAYS = "月火水木金土日"


def norm_headline(value: str) -> str:
    """Compare headlines across Excel and CSV.

    The workbook standardized the old 「▲ = decrease」 convention to ▼.
    CSV exports can still use a bare ▲ for a decrease and ▲+ for an increase.
    """
    text = unicodedata.normalize("NFKC", value or "")
    text = re.sub(r"▲\s*[+＋]", "+", text)
    text = text.replace("▲", "-").replace("▼", "-")
    text = text.replace("＋", "+").replace("−", "-").replace("－", "-")
    return re.sub(r"\s+", "", text)


def category_label(value: str) -> str:
    text = unicodedata.normalize("NFKC", (value or "").strip())
    return re.sub(r"^[^\w]+", "", text, flags=re.UNICODE).strip()


# テーマ判定。
# 見出しを話題ごとに切り、先頭から順に最初に当てはまったテーマを使う。
# 本文は、見出しだけでは決まらない短い要点の補完と、本文先頭の固有名詞に限る。
# メモ（話題の切り口・実務メモ）は判定に使わない。
_SECURITY_RE = re.compile(r"不正アクセス|情報流出|流出|フィッシング|パスポート|個人情報|漏洩|詐欺")
# 「12/1開業」のように日付の直後でも開業とみなす。「開業6年」は対象外。
_OPENING_RE = re.compile(
    r"グランドオープン|起工式|起工|着工|リニューアル開業|リニューアル|新規開業|新ブランド|日本初進出|開業(?!\d)"
)
_OPENING_BLOCK_RE = re.compile(r"営業終了|REIT|RevPAR")
# 成田・羽田だけでなく、関西・福岡・長崎など全国の空港同士の路線（関西＝シアトル など）も対象。
# 新規就航だけでなく、運休・路線廃止・運航再開・チャーター便・運航の移管もフライト。
_FLIGHT_EVENT_RE = re.compile(
    r"就航|増便|減便|新路線|燃油サーチャージ|滑走路|国際線|ダイヤ|臨時便|"
    r"運休|路線廃止|運航再開|チャーター便|移管|"
    r"(?:成田|羽田|関西|関空|伊丹|中部|福岡|新千歳|那覇|広島|仙台|神戸|長崎|熊本|鹿児島|松山|高松|小松|静岡|函館|釧路)[＝=]"
)
_AIRLINE_RE = re.compile(
    r"ZIPAIR|ルフトハンザ|キャセイ|Peach|ユナイテッド|ビーマン|アルキア|エアアスタナ|"
    r"ジェットスター|NZ航空|デルタ|サウディア|ANA|JAL"
)
_POLICY_RE = re.compile(
    r"観光庁|JNTO|国交相|国土交通|宿泊税|応援割|ふっこう|リファンド|宿泊旅行統計|"
    r"訪日外客|訪日消費|インバウンド消費|地方誘客|GO BEYOND|取扱|JESTA|"
    r"オーバーツーリズム|二重価格|国際年|助成|制度調査|旅行業法|消費動向|次速報|"
    r"訪日\d+月|\d+月の?訪日|訪日客は|国・地域別"
)
_DOMESTIC_RE = re.compile(
    r"じゃらん|楽天トラベル|楽天スーパー|楽天「|一休|Relux|JTB|HIS|KNT|近畿日本ツーリスト|"
    r"近ツー|(?<!訪日)日本旅行(?!料金|者|業)|エアトリ|skyticket|スカイチケット"
)
_B2B_EXPLICIT_RE = re.compile(
    r"Hotelbeds|HBX|WebBeds|Webbeds|RateHawk|Ratehawk|Emerging Travel|\bETG\b|"
    r"TAAP|Partner Solutions|Cloudbeds|Virtuoso|Bedsonline|UCP|B2B|商談会"
)
_EXPEDIA_RE = re.compile(r"Expedia|エクスペディア|Vrbo")
# 単独の「OTA」は本文のチャネル統計でも当たるので、見出し判定では使わない。
_OVERSEAS_RE = re.compile(
    r"Booking|Agoda|Trip\.com|トリップドットコム|Airbnb|エアビー|Klook|パリティ|手数料|"
    r"海外OTA|Vayo|TripGenie|Lola|GuestReservations|略奪|HOTREC"
)
_SITE_RE = re.compile(
    r"サイトコントローラー|サイトコントローラ|チャネルマネージャ|Hotel Center|第三者ホテル|コネクティビティ"
)
# 他旅行会社動向。テーマが海外OTA・国内OTAでなくても、見出しが旅行会社・ランドオペ・B2Bなら対象。
# 「HIS」は Hospitality などに部分一致させない。メモは見ない。
_COMPANY_NAME_RE = re.compile(
    r"じゃらん|楽天トラベル|楽天スーパー|楽天「|一休|Relux|"
    r"(?<![A-Za-z0-9])JTB(?![A-Za-z0-9])|"
    r"(?<![A-Za-z])HIS(?![A-Za-z])|エイチ・アイ・エス|"
    r"(?<![A-Za-z])KNT(?![A-Za-z])|近畿日本ツーリスト|近ツー|"
    r"(?<!訪日)日本旅行(?!料金|者|業)|"
    r"エアトリ|skyticket|スカイチケット|"
    r"阪急交通社|阪急トラベル|東武トップツアーズ|クラブツーリズム|農協観光|名鉄観光|"
    r"読売旅行|ビッグホリデー|旅工房|はとバス|ジャルパック|(?<![A-Za-z])JALPAK(?![A-Za-z])|"
    r"ミキ・ツーリスト|Miki\s*Travel|"
    r"ランドオペレータ(?:ー)?|"
    r"ホールセラー|ホールセール|ベッドバンク|[Bb]edbank|"
    r"Hotelbeds|(?<![A-Za-z])HBX(?![A-Za-z])|Web[Bb]eds|Rate[Hh]awk|Emerging Travel|"
    r"(?<![A-Za-z])ETG(?![A-Za-z])|(?<![A-Za-z])TAAP(?![A-Za-z])|Bedsonline|Cloudbeds|Virtuoso|"
    r"(?<![A-Za-z])B2B(?![A-Za-z])|"
    r"Expedia|エクスペディア|(?<![A-Za-z])Vrbo(?![A-Za-z])|"
    r"Booking|Agoda|Trip\.com|トリップドットコム|Airbnb|エアビー|"
    r"Klook|KKday|Veltra|ベルトラ|Viator|Tiqets|GetYourGuide|"
    r"Traveloka|MakeMyTrip|(?<![A-Za-z])Ctrip(?![A-Za-z])|携程|Fliggy|"
    r"サイトコントローラ(?:ー)?|"
    r"(?<![A-Za-z])Dida(?![A-Za-z])|SiteMinder|DerbySoft|手間いらず|ねっぱん|TL-Lincoln|"
    r"大手旅行会社|旅行会社"
)
# 本文から拾うのは会社・規格の固有名だけ（「OTA」「手数料」などの一般語は拾わない）。
_BODY_OVERSEAS_RE = re.compile(
    r"Booking|Agoda|Trip\.com|トリップドットコム|Airbnb|エアビー|Klook|Vayo|TripGenie|Lola"
)


def _norm_text(value: str) -> str:
    return unicodedata.normalize("NFKC", value or "")


def _split_dots(text: str) -> list[str]:
    """中黒で話題を分ける。括弧内の中黒（室数の列挙など）は分割しない。"""
    parts: list[str] = []
    depth = 0
    current: list[str] = []
    for char in text:
        if char in "（(":
            depth += 1
            current.append(char)
        elif char in "）)":
            depth = max(0, depth - 1)
            current.append(char)
        elif char == "・" and depth == 0:
            piece = "".join(current).strip()
            if piece:
                parts.append(piece)
            current = []
        else:
            current.append(char)
    tail = "".join(current).strip()
    if tail:
        parts.append(tail)
    return parts


def _clauses(headline: str, highlight: bool) -> list[str]:
    # NFKC は全角スラッシュを半角にするため、分割してから正規化する。
    raw = headline or ""
    chunks = [chunk.strip() for chunk in re.split(r"[　\s]*／[　\s]*|　+", raw) if chunk.strip()]
    chunks = [_norm_text(chunk) for chunk in chunks]
    text = _norm_text(raw)
    if not highlight:
        return chunks or [text]
    clauses: list[str] = []
    for chunk in chunks:
        clauses.extend(_split_dots(chunk))
    return clauses or [text]


def _first_bullet(body: str) -> str:
    text = _norm_text(body).lstrip("・ 　\n")
    line = re.split(r"[\n。]", text, maxsplit=1)[0]
    return line.strip()[:160]


def _classify(text: str, *, allow_opening: bool) -> str | None:
    """1つの話題テキストをテーマにする。決まらなければ None。"""
    if not text:
        return None
    if _SECURITY_RE.search(text) and not re.search(r"宿泊税|応援割|旅行業法|訪日外客|取扱速報", text):
        return "other"
    if _POLICY_RE.search(text):
        return "policy"
    if allow_opening and _OPENING_RE.search(text) and not _OPENING_BLOCK_RE.search(text):
        return "openings"
    if _FLIGHT_EVENT_RE.search(text) or (
        _AIRLINE_RE.search(text) and re.search(r"就航|増便|減便|ダイヤ|路線|開設|延期|臨時", text)
    ):
        return "flights"
    if re.search(r"REIT|RevPAR", text):
        return "other"
    # Expedia / Hotelbeds / WebBeds / RateHawk / サイトコントローラー / 配信規格は海外OTA。
    if _B2B_EXPLICIT_RE.search(text) or _EXPEDIA_RE.search(text):
        return "overseas"
    if _DOMESTIC_RE.search(text):
        return "domestic"
    if _OVERSEAS_RE.search(text):
        return "overseas"
    # 「OTA」単独はチャネル統計でも出る。AIと同時に出ていれば海外OTAの話題とみなす。
    if re.search(r"(?<![A-Za-z])OTA(?![A-Za-z])", text) and re.search(r"AI|人工知能", text):
        return "overseas"
    if _SITE_RE.search(text):
        return "overseas"
    return None


def is_company_news(item: dict) -> bool:
    """他旅行会社動向に出すか。

    海外OTAと国内OTA・旅行会社は常に対象。それ以外は見出し・カテゴリ、
    または本文の書き出しが旅行会社・ランドオペレーター・卸/B2Bのときだけ。
    開業記事の本文で予約先としてOTA名が並ぶだけでは含めない。
    note（話題の切り口・実務メモ）は判定に使わない。
    """
    if item.get("theme") in ("overseas", "domestic"):
        return True
    text = _norm_text(
        "\n".join(
            [
                item.get("headline") or "",
                item.get("category") or "",
                item.get("categoryLabel") or "",
            ]
        )
    )
    if _COMPANY_NAME_RE.search(text):
        return True
    lead = _first_bullet(item.get("body") or "")
    match = _COMPANY_NAME_RE.search(lead)
    return bool(match and match.start() <= 2)


def assign_theme(category: str, headline: str, body: str, note: str) -> str:
    """項目をちょうど1つのテーマに振り分ける。どれにも乗らなければ other。

    note は打合せ用メモなので判定に使わない（将来の行でも同じ）。
    """
    del note  # 判定対象外であることを呼び出し側に残す
    label = category_label(category)
    head = _norm_text(headline)
    highlight = label == "要点"

    # 「東京国際線」に限らず、シートの「国際線」（関西・長崎など全国の路線）もフライト。
    if label == "国際線" or label.endswith("国際線"):
        return "flights"
    if "宿泊税" in label or label in ("政策・支援", "国内政策"):
        return "policy"
    if "セキュリティ" in label or label.startswith("システム"):
        return "other"

    # 生の見出しを渡す。NFKC すると話題区切りの「／」が半角スラッシュになり、日付と区別できなくなる。
    clauses = _clauses(headline, highlight)
    first_bullet = _first_bullet(body)
    if highlight and clauses:
        # 要点は先頭の話題を優先する。見出しだけでは決まらないときだけ、最初の箇条書きを足す。
        lead = _classify(clauses[0], allow_opening=True)
        if not lead:
            lead = _classify(clauses[0] + "\n" + first_bullet, allow_opening=True)
        if lead:
            return lead
        for clause in clauses[1:]:
            theme = _classify(clause, allow_opening=True)
            if theme:
                return theme
    else:
        for clause in clauses:
            theme = _classify(clause, allow_opening=True)
            if theme:
                return theme

    # 見出しが会社名も制度名も持たない短い要点は、本文の書き出しで補う。
    if highlight and len(head) <= 28:
        hinted = _classify(head + "\n" + first_bullet, allow_opening=True)
        if hinted:
            return hinted

    # JNTO・観光庁の発表もの（受賞、統計の2次速報、補助の公募延長など）。
    if re.match(r"^(?:JNTO|観光庁)", first_bullet):
        return "policy"

    # 見出しにテーマが無いときは、本文の最初の文にある会社名だけを見る。
    if not highlight:
        if _B2B_EXPLICIT_RE.search(first_bullet) or _EXPEDIA_RE.search(first_bullet) or _SITE_RE.search(first_bullet):
            return "overseas"
        if _DOMESTIC_RE.search(first_bullet):
            return "domestic"
        if _BODY_OVERSEAS_RE.search(first_bullet):
            return "overseas"

    if label == "ホテル":
        return "openings" if _classify(head, allow_opening=True) == "openings" else "other"
    if "国内OTA" in label or label in ("旅行会社", "国内・JTB", "国内販促", "国内"):
        return "policy" if re.search(r"応援割|ふっこう|宿泊税", head) else "domestic"
    if "Hotelbeds" in label or label in ("配信・規格",):
        return "overseas"
    if any(key in label for key in ("海外OTA", "Agoda", "Airbnb", "Booking", "Trip.com", "Klook", "規制", "パリティ", "手数料", "Expedia")):
        return "overseas"
    return "other"


def cell_text(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%m/%d")
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def parse_month_day(value: str) -> tuple[int, int] | None:
    match = re.match(r"\s*(\d{1,2})\s*/\s*(\d{1,2})", cell_text(value))
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def weekday_from_label(value: str) -> str:
    match = re.search(r"[（(]([月火水木金土日])[）)]", cell_text(value))
    return match.group(1) if match else ""


def split_urls(text: str) -> list[str]:
    if not text:
        return []
    found = re.findall(r"https?://[^\s;<>]+", text.strip().strip('"'))
    seen: set[str] = set()
    urls: list[str] = []
    for raw in found:
        url = raw.rstrip(").,、」』")
        if url not in seen:
            seen.add(url)
            urls.append(url)
    return urls


def host_label(url: str) -> str:
    host = urlparse(url).netloc.lower()
    if host.startswith("www."):
        host = host[4:]
    return host or url


def load_csv(path: Path | None) -> list[dict]:
    if path is None:
        return []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def find_header(ws):
    for index, row in enumerate(ws.iter_rows(values_only=True), 1):
        values = [cell_text(cell) for cell in row]
        if "日付" in values:
            return index, {name: pos for pos, name in enumerate(values) if name}
    raise SystemExit(f"{ws.title}: ヘッダー行（日付）が見つかりません")


def read_sheet(ws, spec: dict) -> list[dict]:
    header_row, columns = find_header(ws)
    required = ["日付", "カテゴリ", "見出し", spec["body"], spec["note"], spec["sources"]]
    missing = [name for name in required if name not in columns]
    if missing:
        raise SystemExit(f"{ws.title}: 列がありません: {', '.join(missing)}")

    items = []
    for row in ws.iter_rows(min_row=header_row + 1):
        date_cell = row[columns["日付"]]
        date_raw = cell_text(date_cell.value)
        if not date_raw:
            continue
        month_day = parse_month_day(date_raw)
        if month_day is None:
            print(f"warning: {ws.title} 日付を解釈できません: {date_raw!r}", file=sys.stderr)
            continue
        source_cell = row[columns[spec["sources"]]]
        hyperlink = source_cell.hyperlink.target if source_cell.hyperlink else ""
        category = cell_text(row[columns["カテゴリ"]].value)
        label = category_label(category)
        headline = cell_text(row[columns["見出し"]].value)
        body = cell_text(row[columns[spec["body"]]].value)
        note = cell_text(row[columns[spec["note"]]].value)
        theme = assign_theme(category, headline, body, note)
        items.append(
            {
                "feed": spec["feed"],
                "feedLabel": spec["feedLabel"],
                "dateRaw": date_raw,
                "month": month_day[0],
                "day": month_day[1],
                "weekday": weekday_from_label(date_raw),
                "category": category,
                "categoryLabel": label,
                "theme": theme,
                "themeLabel": THEME_LABEL[theme],
                "highlight": label == "要点",
                "headline": headline,
                "body": body,
                "note": note,
                "noteLabel": spec["noteLabel"],
                "excelLink": hyperlink if hyperlink and hyperlink.startswith("http") else "",
                "sourceText": cell_text(source_cell.value),
            }
        )
    return items


def index_csv(rows: list[dict], source_key: str) -> dict:
    """Index CSV rows by (month, day, normalized headline)."""
    indexed: dict[tuple, list] = {}
    for row in rows:
        iso = (row.get("日付") or "").strip()
        match = re.match(r"(\d{4})-(\d{2})-(\d{2})", iso)
        if not match:
            continue
        headline = row.get("見出し") or ""
        key = (int(match.group(2)), int(match.group(3)), norm_headline(headline))
        indexed.setdefault(key, []).append(
            {
                "iso": f"{match.group(1)}-{match.group(2)}-{match.group(3)}",
                "urls": split_urls(row.get(source_key) or ""),
                "category": category_label(row.get("カテゴリ") or ""),
            }
        )
    return indexed


def assign_dates_and_sources(items: list[dict], csv_index: dict, year: int) -> None:
    """Fill ISO dates and source URLs. Dates walk newest-first within the sheet."""
    previous: tuple[int, int] | None = None
    current_year = year
    for item in items:
        key = (item["month"], item["day"], norm_headline(item["headline"]))
        hits = list(csv_index.get(key, []))
        if len(hits) > 1:
            same_cat = [hit for hit in hits if hit["category"] == item["categoryLabel"]]
            if same_cat:
                hits = same_cat
        chosen = hits[0] if hits else None

        if chosen:
            item["date"] = chosen["iso"]
            current_year = int(chosen["iso"][:4])
            urls = list(chosen["urls"])
            item["sourceMatch"] = "csv"
        else:
            if previous is not None and (item["month"], item["day"]) > previous:
                current_year -= 1
            item["date"] = f"{current_year:04d}-{item['month']:02d}-{item['day']:02d}"
            urls = []
            item["sourceMatch"] = "excel" if item["excelLink"] else "none"

        if not urls and item["excelLink"]:
            urls = [item["excelLink"]]
            if item["sourceMatch"] == "csv":
                item["sourceMatch"] = "excel-fallback"
        elif item["excelLink"] and item["excelLink"] not in urls:
            urls.insert(0, item["excelLink"])

        if not item["weekday"]:
            # Monday=0 in Python; Japanese week starts on Monday with the same index.
            import datetime

            item["weekday"] = WEEKDAYS[datetime.date.fromisoformat(item["date"]).weekday()]

        item["sources"] = [{"url": url, "label": host_label(url)} for url in urls]
        previous = (item["month"], item["day"])


def public_item(item: dict, seq: int) -> dict:
    return {
        "id": f"{item['feed']}-{seq}",
        "date": item["date"],
        "dateLabel": f"{int(item['date'][5:7])}/{int(item['date'][8:10])}（{item['weekday']}）",
        "weekday": item["weekday"],
        "feed": item["feed"],
        "feedLabel": item["feedLabel"],
        "category": item["category"],
        "categoryLabel": item["categoryLabel"],
        "theme": item["theme"],
        "themeLabel": item["themeLabel"],
        "highlight": item["highlight"],
        "headline": item["headline"],
        "body": item["body"],
        "note": item["note"],
        "noteLabel": item["noteLabel"],
        "sources": item["sources"],
        "companies": is_company_news(item),
    }


def build(workbook_path: Path, industry_csv: Path | None, ota_csv: Path | None, year: int) -> dict:
    wb = openpyxl.load_workbook(workbook_path, data_only=False, read_only=False)
    missing = [name for name in SHEETS if name not in wb.sheetnames]
    if missing:
        raise SystemExit(
            f"シートが見つかりません: {', '.join(missing)}（ファイル内: {', '.join(wb.sheetnames)}）"
        )

    industry_items = read_sheet(wb[INDUSTRY_SHEET], SHEETS[INDUSTRY_SHEET])
    ota_items = read_sheet(wb[OTA_SHEET], SHEETS[OTA_SHEET])
    assign_dates_and_sources(
        industry_items,
        index_csv(load_csv(industry_csv), "出典"),
        year,
    )
    assign_dates_and_sources(
        ota_items,
        index_csv(load_csv(ota_csv), "ソース"),
        year,
    )

    combined = industry_items + ota_items
    # Newest date first. Within a date, 要点 first, then the workbook order
    # (industry sheet, then OTA sheet). Sort is stable, so equal keys keep that order.
    # reverse=True applies to the whole key, so the highlight flag is 1 for 要点.
    combined.sort(key=lambda item: (item["date"], 1 if item["highlight"] else 0), reverse=True)
    omitted = [item for item in combined if item["theme"] in OMITTED_IDS]
    combined = [item for item in combined if item["theme"] not in OMITTED_IDS]

    items = [public_item(item, index + 1) for index, item in enumerate(combined, 1)]
    theme_counts = {theme_id: sum(1 for item in items if item["theme"] == theme_id) for theme_id, _label in THEMES}
    filters = [{"id": "all", "label": "すべて"}]
    for theme_id, label in CHIP_THEMES:
        filters.append({"id": theme_id, "label": label})
    updated = max((item["date"] for item in items), default="")
    return {
        "updated": updated,
        "year": year,
        "source": workbook_path.name,
        "filters": filters,
        "counts": {
            "items": len(items),
            "dates": len({item["date"] for item in items}),
            "industry": sum(1 for item in items if item["feed"] == "industry"),
            "ota": sum(1 for item in items if item["feed"] == "ota"),
            "industryDates": len({item["date"] for item in items if item["feed"] == "industry"}),
            "otaDates": len({item["date"] for item in items if item["feed"] == "ota"}),
            "themes": {THEME_LABEL[theme_id]: theme_counts[theme_id] for theme_id, _label in THEMES},
        },
        "items": items,
        "_audit": {
            "industrySourceMatch": _tally(industry_items),
            "otaSourceMatch": _tally(ota_items),
            "omitted": {label: sum(1 for item in omitted if item["theme"] == theme_id) for theme_id, label in OMITTED_THEMES},
        },
    }


def _tally(items: list[dict]) -> dict:
    tally: dict[str, int] = {}
    for item in items:
        tally[item["sourceMatch"]] = tally.get(item["sourceMatch"], 0) + 1
    return tally


def main() -> None:
    parser = argparse.ArgumentParser(description="旅行業界・海外OTAニュースの Excel から news.json を作ります。")
    parser.add_argument("xlsx", type=Path, help="マスター Excel（旅行業界ニュース / 海外OTAニュース）")
    parser.add_argument("--industry-csv", type=Path, default=None, help="旅行業界シートの CSV（出典の全URL）")
    parser.add_argument("--ota-csv", type=Path, default=None, help="海外OTAシートの CSV（ソースの全URL）")
    parser.add_argument("-o", "--output", type=Path, default=Path("news.json"))
    parser.add_argument("--year", type=int, default=2026, help="日付に年が無い行の年（既定 2026）。新しい順で月が戻ったら前年とみなす")
    args = parser.parse_args()

    if not args.xlsx.is_file():
        raise SystemExit(f"ファイルがありません: {args.xlsx}")

    payload = build(args.xlsx, args.industry_csv, args.ota_csv, args.year)
    audit = payload.pop("_audit")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    counts = payload["counts"]
    print(f"wrote {args.output}  items={counts['items']}  dates={counts['dates']}  updated={payload['updated']}")
    print(
        f"  旅行業界: {counts['industry']} items / {counts['industryDates']} dates"
        f"  sources={audit['industrySourceMatch']}"
    )
    print(
        f"  海外OTA: {counts['ota']} items / {counts['otaDates']} dates"
        f"  sources={audit['otaSourceMatch']}"
    )
    unmatched = []
    for item in payload["items"]:
        if not item["headline"] or not item["body"]:
            unmatched.append(f"empty text {item['id']} {item['date']} {item['headline'][:40]}")
        if not item["sources"]:
            unmatched.append(f"no sources {item['id']} {item['date']} {item['categoryLabel']} {item['headline'][:40]}")
    print("  テーマ別:")
    for theme_id, label in THEMES:
        print(f"    {label}: {counts['themes'][label]}")
    print(f"  他旅行会社動向: {sum(1 for item in payload['items'] if item.get('companies'))}")
    omitted = audit.get("omitted") or {}
    if any(omitted.values()):
        print(
            "  一覧から除外: "
            + "、".join(f"{theme_id} {n}" for theme_id, n in omitted.items())
            + f"（計 {sum(omitted.values())} 件）"
        )
    if unmatched:
        print(f"  要確認 {len(unmatched)} 件:")
        for line in unmatched:
            print(f"    {line}")


if __name__ == "__main__":
    main()
